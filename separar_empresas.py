#!/usr/bin/env python3
"""Separa a planilha em um arquivo por empresa.

Uso:
    python separar_empresas.py planilha.xlsx [-o saida] [--cabecalho 1]

Para cada empresa distinta da coluna D da aba "Análise Previdenciária Geral",
copia o arquivo inteiro e apaga as linhas das demais empresas, salvando como
"<EMPRESA> - Análise Previdenciária.xlsx".
"""
import argparse
import re
import shutil
from pathlib import Path

from openpyxl import load_workbook

ABA = "Análise Previdenciária Geral"
COL_EMPRESA = 4  # coluna D


def nome_seguro(nome: str) -> str:
    return re.sub(r'[\\/:*?"<>|]', "-", nome).strip().rstrip(".")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("planilha")
    ap.add_argument("-o", "--saida", default="saida")
    ap.add_argument("--aba", default=ABA)
    ap.add_argument("--cabecalho", type=int, default=1,
                    help="linhas de cabeçalho a preservar (padrão: 1)")
    args = ap.parse_args()

    origem = Path(args.planilha)
    saida = Path(args.saida)
    saida.mkdir(parents=True, exist_ok=True)

    ws = load_workbook(origem)[args.aba]
    empresas = []
    for r in range(args.cabecalho + 1, ws.max_row + 1):
        v = ws.cell(r, COL_EMPRESA).value
        if v is not None and str(v).strip() and str(v).strip() not in empresas:
            empresas.append(str(v).strip())
    print(f"{len(empresas)} empresas encontradas: {empresas}")

    for emp in empresas:
        destino = saida / f"{nome_seguro(emp)} - Análise Previdenciária.xlsx"
        shutil.copy(origem, destino)
        wb = load_workbook(destino)
        ws = wb[args.aba]
        # de baixo para cima, para os índices não se deslocarem
        for r in range(ws.max_row, args.cabecalho, -1):
            v = ws.cell(r, COL_EMPRESA).value
            if str(v).strip() != emp if v is not None else True:
                ws.delete_rows(r)
        wb.save(destino)
        print(f"  gerado: {destino}")


if __name__ == "__main__":
    main()
